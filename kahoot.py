"""Kahoot natijalarini o'qish va talaba ismlariga moslash."""
import difflib
import re
import unicodedata

import openpyxl

APOSTROPHES = "'`ʻʼ‘’ʹ"


def normalize(name: str) -> str:
    """Ismni solishtirish uchun soddalashtiradi: kichik harf, tutuq belgisiz."""
    s = unicodedata.normalize("NFKC", str(name or "")).lower()
    for ch in APOSTROPHES:
        s = s.replace(ch, "")
    s = re.sub(r"[^a-zа-яёқғўҳ\s]", " ", s)
    # Otasining ismi qo'shimchalari moslashga xalaqit bermasin
    s = re.sub(r"\b(qizi|ogli|o g li|ugli|kizi)\b", " ", s)
    return " ".join(s.split())


def _tokens(name: str) -> set:
    return set(normalize(name).split())


def _to_int(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def parse_kahoot_xlsx(path: str):
    """Kahoot'dan yuklangan Excel hisobotni o'qiydi.

    Qaytaradi: [{"name", "correct", "total"}] yoki foizli hisobotda [{"name", "percent"}], izoh
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    # Savollar soni: nomi raqam bilan boshlanadigan varaqlar (har savolga bitta varaq)
    q_sheets = [ws for ws in wb.sheetnames if re.match(r"^\s*\d+", ws)]
    n_questions = len(q_sheets) or None

    target = None
    for name in wb.sheetnames:
        if "final" in name.lower() or "score" in name.lower():
            target = wb[name]
            break
    sheets = [target] if target else [wb[n] for n in wb.sheetnames]

    for ws in sheets:
        rows = list(ws.iter_rows(values_only=True))
        for i, row in enumerate(rows[:30]):
            cells = [str(c).strip().lower() if c is not None else "" for c in row]
            name_col = next((j for j, c in enumerate(cells) if c in ("player", "players", "nickname", "name")), None)
            corr_col = next((j for j, c in enumerate(cells) if c.startswith("correct")), None)
            if name_col is None or corr_col is None:
                continue
            inc_col = next((j for j, c in enumerate(cells) if c.startswith("incorrect")), None)
            una_col = next((j for j, c in enumerate(cells) if c.startswith("unanswered") or c.startswith("no answer")), None)
            body = [r for r in rows[i + 1:] if r and name_col < len(r) and r[name_col] not in (None, "")]
            # Topshiriq (assignment) hisobotida "Correct answers" foizda bo'ladi: 74% yoki 0.74
            raw = [r[corr_col] for r in body if corr_col < len(r)]
            percent_mode = "%" in cells[corr_col] or any(
                (isinstance(v, str) and "%" in v) or (isinstance(v, float) and 0 < v < 1 and v != int(v)) for v in raw)
            out = []
            for r in body:
                v = r[corr_col] if corr_col < len(r) else None
                name = str(r[name_col]).strip()
                if percent_mode:
                    p = _to_percent(v)
                    if p is not None:
                        out.append({"name": name, "percent": p})
                    continue
                correct = _to_int(v)
                if correct is None:
                    continue
                total = n_questions
                if total is None:
                    total = correct + ((_to_int(r[inc_col]) or 0) if inc_col is not None else 0)
                    if una_col is not None:
                        total += _to_int(r[una_col]) or 0
                out.append({"name": name, "correct": correct, "total": total})
            if out:
                if percent_mode:
                    return out, f"{len(out)} ta o'yinchi (foizli hisobot)"
                # Savollar soni noma'lum bo'lsa, eng ko'p javob bergan talabaga qarab olamiz
                if n_questions is None:
                    mx = max(x["total"] for x in out)
                    for x in out:
                        x["total"] = mx
                return out, f"{len(out)} ta o'yinchi, {out[0]['total']} ta savol"
    raise ValueError("Excel faylda Kahoot natijalari jadvali (Player / Correct answers) topilmadi")


def _to_percent(v):
    if v is None or v == "":
        return None
    if isinstance(v, str):
        try:
            return float(v.replace("%", "").replace(",", ".").strip())
        except ValueError:
            return None
    v = float(v)
    return v * 100 if v <= 1 else v


def score10(res) -> float:
    """Natijadan 0-10 ball: foiz bo'lsa foiz/10, aks holda to'g'ri/jami x 10."""
    if res.get("percent") is not None:
        return round(min(max(float(res["percent"]), 0), 100) / 10, 1)
    if not res.get("total"):
        return 0.0
    return round(res["correct"] / res["total"] * 10, 1)


def match_students(results, students, aliases=None, cutoff=0.82):
    """Kahoot taxalluslarini jurnaldagi ismlarga bog'laydi.

    students: {qator_raqami: "To'liq ism"}
    aliases: {normalize(taxallus): "To'liq ism"}
    Qaytaradi: matched [(qator, ism, natija)], unmatched [natija]
    """
    aliases = aliases or {}
    by_norm = {normalize(n): row for row, n in students.items()}
    tok = {row: _tokens(n) for row, n in students.items()}
    matched, unmatched, used = [], [], set()

    for res in results:
        nick = normalize(res["name"])
        row = None
        if nick in aliases and normalize(aliases[nick]) in by_norm:
            row = by_norm[normalize(aliases[nick])]
        if row is None and nick in by_norm:
            row = by_norm[nick]
        if row is None and nick:
            nt = set(nick.split())
            # Taxallusdagi barcha so'zlar faqat bitta talabaning ismida bo'lsa
            def covers(t):
                # "Tursunova S" ham "Tursunova Sevara" ga mos keladi
                return all(w in t or (len(w) <= 2 and any(x.startswith(w) for x in t)) for w in nt)
            # Kamida bitta to'liq so'z bo'lishi shart: "B" yoki "Sh" yolg'iz o'zi yetmaydi
            strong = any(len(w) >= 3 for w in nt)
            subset = [r for r, t in tok.items() if nt and strong and covers(t)]
            if len(subset) == 1:
                row = subset[0]
            else:
                key = " ".join(sorted(nt))
                scored = sorted(
                    ((difflib.SequenceMatcher(None, key, " ".join(sorted(t))).ratio(), r) for r, t in tok.items()),
                    reverse=True,
                )
                if scored and scored[0][0] >= cutoff and (len(scored) == 1 or scored[0][0] - scored[1][0] > 0.05):
                    row = scored[0][1]
        if row is None or row in used:
            unmatched.append(res)
        else:
            used.add(row)
            matched.append((row, students[row], res))
    return matched, unmatched
